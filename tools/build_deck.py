"""Builds deck/Board_Strategic_Plan_FY2027.pptx: a board-style summary of the plan.

All numbers come from reference_model.py, which verify_workbook.py proves equal to the workbook.
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import assumptions as A
import reference_model as R

OUT = Path(__file__).resolve().parents[1] / "deck" / "Board_Strategic_Plan_FY2027.pptx"
NAVY, GREY, RED, GREEN = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x59, 0x59, 0x59), RGBColor(0xC0, 0x00, 0x00), RGBColor(0x37, 0x86, 0x3C)
TEAL, GOLD, LIGHT = RGBColor(0x2E, 0x75, 0xB6), RGBColor(0xBF, 0x8F, 0x00), RGBColor(0xD9, 0xE1, 0xF2)
SERIES = [NAVY, TEAL, GOLD, RGBColor(0x7F, 0x7F, 0x7F)]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=16, color=GREY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.color.rgb, p.font.bold, p.font.name = Pt(size), color, bold, "Calibri"
        p.space_after = Pt(6)
        if align:
            p.alignment = align
    return box


def slide(title, subtitle=None, n=[1]):
    n[0] += 1
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text(s, 0.5, 0.3, 12.3, 0.8, title, 28, NAVY, True)
    if subtitle:
        text(s, 0.5, 0.95, 12.3, 0.5, subtitle, 15, GREY)
    text(s, 0.5, 7.0, 9, 0.3, f"{A.COMPANY} (illustrative) | FY2027 Budget & Strategic Plan | $ millions", 10, GREY)
    text(s, 12.2, 7.0, 0.7, 0.3, str(n[0]), 10, GREY, align=PP_ALIGN.RIGHT)
    return s


def table(s, x, y, w, rows, col_w=None, size=12, red_cells=(), left_cols=(0,)):
    shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.36 * len(rows)))
    t = shape.table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(v)
            p = cell.text_frame.paragraphs[0]
            p.font.size, p.font.name = Pt(size), "Calibri"
            p.font.bold = i == 0
            p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else (RED if (i, j) in red_cells else GREY)
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.RIGHT
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    return t


def style_chart(chart, legend=True, size=12):
    chart.font.size, chart.font.name = Pt(size), "Calibri"
    chart.has_legend = legend
    if legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False


m = lambda v: f"{v:,.0f}"
pct = lambda v: f"{v * 100:.1f}%"

c, plans = R.consolidation(1)
scen = {k: R.summary(k) for k in A.SCENARIOS}
base = scen[1]
bridge = R.ebitda_bridge(1)
proj = R.growth_projects()
cap26 = R.capex_2026()
hist, ytd, latest = A.macro_annual()
below = proj[proj.npv < 0]

# 1. Title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, 0.8, 2.3, 11.5, 1.2, "FY2027 Budget & Five-Year Strategic Plan", 40, RGBColor(0xFF, 0xFF, 0xFF), True)
text(s, 0.8, 3.4, 11.5, 0.6, f"{A.COMPANY}  |  Board of Directors review", 22, RGBColor(0xFF, 0xFF, 0xFF))
text(s, 0.8, 5.8, 11.5, 0.8, ["Illustrative company and figures for a portfolio project. Macro data: FRED, Federal Reserve Bank of St. Louis.",
                              "Prepared by Kshama Upadhyay"], 13, RGBColor(0xD9, 0xE1, 0xF2))

# 2. Executive summary
s = slide("Executive summary", "Base case; FY2026 is the latest 9+3 forecast")
text(s, 0.6, 1.6, 12.2, 5.2, [
    f"FY2027 budget: revenue ${m(c.loc[2027, 'revenue'])}M (+{pct(c.loc[2027, 'revenue'] / c.loc[2026, 'revenue'] - 1)}) and EBITDA ${m(c.loc[2027, 'ebitda'])}M "
    f"(+{pct(c.loc[2027, 'ebitda'] / c.loc[2026, 'ebitda'] - 1)}), margin {pct(c.loc[2027, 'margin'])}.",
    f"Five-year plan: revenue grows {pct(base['revenue_cagr'])} a year to ${m(base['revenue_2031'])}M and EBITDA {pct(base['ebitda_cagr'])} a year to "
    f"${m(base['ebitda_2031'])}M by FY2031; margin expands {(base['margin_2031'] - c.loc[2026, 'margin']) * 100:.1f} pts.",
    f"Price is the main lever: +${m(bridge['price'])}M of the +${m(bridge['end'] - bridge['start'])}M EBITDA increase, offset by "
    f"${m(-bridge['var_cost'] - bridge['fixed'])}M of cost inflation. Volume adds ${m(bridge['volume'])}M, growth projects ${m(bridge['growth'])}M.",
    f"Capital: ${m(base['capex_5yr'])}M of capex over FY2027-FY2031 (${m(base['growth_capex_5yr'])}M growth); "
    f"cumulative free cash flow ${m(base['fcf_5yr'])}M, {pct(base['fcf_conversion_5yr'])} of EBITDA.",
    f"FY2026 capex is forecast ${m(cap26.variance.sum())}M under budget, mainly from the Southeast quarry permitting delay; kiln project is $6M over approval.",
    f"{len(below)} growth projects ({', '.join(below.index)}) fall below the {pct(A.WACC)} hurdle rate; we recommend re-scoping them before release of funds.",
    f"Downside (housing and highway volumes -4/-3 pts in FY2027-28, price -1 pt): FY2031 EBITDA ${m(scen[2]['ebitda_2031'])}M, "
    f"still ${m(scen[2]['fcf_5yr'])}M of five-year free cash flow.",
], 17)

# 3. Macro backdrop
s = slide("Macro: construction demand is flat, highway spending still growing",
          "Annual averages, indexed to 2019 = 100 (2026 = year to date). Source: FRED")
cd = CategoryChartData()
years = list(hist.index)
cd.categories = [str(y) if y < 2026 else "2026 YTD" for y in years]
for sid, label in [("HOUST", "Housing starts"), ("TTLCONS", "Total construction spending"),
                   ("TLHWYCONS", "Highway & street spending"), ("WPUSI012011", "PPI: construction materials")]:
    cd.add_series(label, [round(hist.loc[y, sid] / hist.loc[2019, sid] * 100, 1) for y in years])
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(8.2), Inches(5.3), cd).chart
style_chart(ch)
for i, ser in enumerate(ch.series):
    ser.format.line.color.rgb = SERIES[i]
    ser.format.line.width = Pt(2.5)
    ser.smooth = False
text(s, 9.0, 1.6, 3.9, 5.2, [
    "2026 year to date vs. 2025:",
    f"Housing starts {pct(ytd['HOUST'])}",
    f"Total construction spending {pct(ytd['TTLCONS'])}",
    f"Highway & street spending +{pct(ytd['TLHWYCONS'])}",
    f"Construction materials prices +{pct(ytd['WPUSI012011'])}",
    f"30-year mortgage rate {latest['MORTGAGE30US'][1]:.2f}% ({latest['MORTGAGE30US'][0]})",
    "Plan assumes housing recovers from FY2028 as rates ease; highway growth slows to 2% from FY2029.",
], 15)

# 4. Five-year plan
s = slide("Five-year plan: EBITDA grows from $%s M to $%s M" % (m(c.loc[2026, "ebitda"]), m(c.loc[2031, "ebitda"])),
          "EBITDA by operating company before corporate costs; base case")
cd = CategoryChartData()
plan_years = list(range(2026, 2032))
cd.categories = ["FY2026F", "FY2027B"] + [f"FY{y}" for y in range(2028, 2032)]
for n, df in plans.items():
    cd.add_series(n, [round(df.loc[y, "ebitda"], 1) for y in plan_years])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(1.5), Inches(7.4), Inches(5.3), cd).chart
style_chart(ch)
for i, ser in enumerate(ch.series):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = SERIES[i]
rows = [["$M", "FY2026F", "FY2027B", "FY2031"]]
for label, key, f in [("Revenue", "revenue", m), ("EBITDA", "ebitda", m), ("EBITDA margin", "margin", pct),
                      ("Capex", "capex", m), ("Free cash flow", "fcf", m)]:
    rows.append([label] + [f(c.loc[y, key]) for y in (2026, 2027, 2031)])
table(s, 8.2, 1.6, 4.7, rows, [1.7, 1.0, 1.0, 1.0], 13)
text(s, 8.2, 4.1, 4.7, 2.5, [f"Aggregates carries the plan: {pct(plans['Aggregates'].loc[2031, 'ebitda'] / c.loc[2031, 'opco_ebitda'])} of FY2031 "
                             f"operating EBITDA at a {pct(plans['Aggregates'].loc[2031, 'margin'])} margin.",
                             f"Intercompany sales of ${m(-c.loc[2027, 'ic_elim'])}M are eliminated in the FY2027 budget."], 14)

# 5. EBITDA bridge (waterfall built from a stacked bar with an invisible base)
s = slide("EBITDA bridge FY2026F to FY2031: price more than covers cost inflation", "Base case, $ millions")
steps = [("FY2026F", bridge["start"]), ("Volume", bridge["volume"]), ("Price", bridge["price"]),
         ("Variable cost\ninflation", bridge["var_cost"]), ("Fixed cost\ninflation", bridge["fixed"]),
         ("Growth\nprojects", bridge["growth"]), ("Corporate", bridge["corporate"]), ("FY2031", bridge["end"])]
base_vals, up, down, total, run = [], [], [], [], 0.0
for i, (label, v) in enumerate(steps):
    if i in (0, len(steps) - 1):
        base_vals.append(0); up.append(0); down.append(0); total.append(round(v, 1)); run = v
    else:
        base_vals.append(round(run if v >= 0 else run + v, 1))
        up.append(round(v, 1) if v >= 0 else 0); down.append(round(-v, 1) if v < 0 else 0); total.append(0)
        run += v
cd = CategoryChartData()
cd.categories = [f"{x[0]}\n{x[1]:+,.0f}" if 0 < i < len(steps) - 1 else f"{x[0]}\n{x[1]:,.0f}" for i, x in enumerate(steps)]
for label, vals in [("base", base_vals), ("Total", total), ("Increase", up), ("Decrease", down)]:
    cd.add_series(label, vals)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3), cd).chart
style_chart(ch, legend=False)
ch.plots[0].gap_width = 60
for ser, color in zip(ch.series, [None, NAVY, GREEN, RED]):
    if color is None:
        ser.format.fill.background()
    else:
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = color

# 6. Operating companies
s = slide("Operating companies: aggregates and cement drive margin expansion")
rows = [["Operating company", "Volume driver", "Revenue FY26F", "Revenue FY31", "CAGR", "Margin FY26F", "Margin FY31"]]
for o in A.OPCOS:
    df = plans[o["name"]]
    rows.append([o["name"], A.MACRO_SERIES[o["driver"]].split(" (")[0], m(df.loc[2026, "revenue"]), m(df.loc[2031, "revenue"]),
                 pct((df.loc[2031, "revenue"] / df.loc[2026, "revenue"]) ** 0.2 - 1), pct(df.loc[2026, "margin"]), pct(df.loc[2031, "margin"])])
rows.append(["Intercompany eliminations", "", m(c.loc[2026, "ic_elim"]), m(c.loc[2031, "ic_elim"]), "", "", ""])
rows.append(["Group (after corporate costs)", "", m(c.loc[2026, "revenue"]), m(c.loc[2031, "revenue"]),
             pct(base["revenue_cagr"]), pct(c.loc[2026, "margin"]), pct(c.loc[2031, "margin"])])
table(s, 0.5, 1.5, 12.3, rows, [3.0, 3.3, 1.2, 1.2, 1.1, 1.25, 1.25], 14, left_cols=(0, 1))
text(s, 0.5, 4.8, 12.3, 2.0, [
    "Volumes follow their end market: aggregates and asphalt track highway spending, cement tracks total construction, ready-mix tracks housing starts.",
    "Ready-mix and asphalt margins stay flat: price increases only keep pace with cement, aggregates and energy costs.",
], 15)

# 7. Capex
s = slide("CAPEX: FY2026 tracking under budget; $%s M five-year program" % m(base["capex_5yr"]))
cd = CategoryChartData()
cd.categories = list(cap26.index)
cd.add_series("FY2026 budget", [round(v, 1) for v in cap26.budget])
cd.add_series("Year to date (Sep)", [round(v, 1) for v in cap26.ytd])
cd.add_series("Full-year forecast", [round(v, 1) for v in cap26.forecast])
text(s, 0.5, 1.25, 6.3, 0.4, "FY2026 capex by operating company", 14, NAVY, True)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.65), Inches(6.3), Inches(4.0), cd).chart
style_chart(ch, size=11)
for i, ser in enumerate(ch.series):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = [LIGHT, TEAL, NAVY][i]
cd = CategoryChartData()
cd.categories = [f"FY{y}" for y in range(2027, 2032)]
cd.add_series("Maintenance", [round(c.loc[y, "maint_capex"], 1) for y in range(2027, 2032)])
cd.add_series("Growth", [round(c.loc[y, "growth_capex"], 1) for y in range(2027, 2032)])
text(s, 7.0, 1.25, 5.9, 0.4, "Five-year capex plan: maintenance vs. growth", 14, NAVY, True)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(7.0), Inches(1.65), Inches(5.9), Inches(4.0), cd).chart
style_chart(ch, size=11)
for i, ser in enumerate(ch.series):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = [NAVY, GOLD][i]
flagged = cap26[cap26.flag != ""]
text(s, 0.5, 5.75, 12.3, 1.2, [
    f"FY2026: budget ${m(cap26.budget.sum())}M, spent ${m(cap26.ytd.sum())}M through September, forecast ${m(cap26.forecast.sum())}M "
    f"(${m(cap26.variance.sum())}M under). Flagged: " + (", ".join(f"{n} {pct(r.var_pct)} under" for n, r in flagged.iterrows()) or "none") + ".",
    f"FY2027 peaks at ${m(c.loc[2027, 'capex'])}M as the quarry, kiln and terminal projects overlap; capex runs "
    f"{c.loc[2027:2031, 'capex'].sum() / c.loc[2027:2031, 'da'].sum():.2f}x D&A over the plan.",
], 14)

# 8. Growth projects
s = slide("Growth projects: six clear the hurdle rate, two need re-scoping",
          f"After-tax cash flows over 20 years, NPV at {pct(A.WACC)} WACC")
rows = [["ID", "Project", "Operating company", "Forecast cost", "vs. approved", "In service", "NPV", "IRR", "Payback (yrs)", "Status"]]
red = set()
for i, (pid, p) in enumerate(proj.sort_values("npv", ascending=False).iterrows(), start=1):
    rows.append([pid, p["name"], p.opco, m(p.forecast_total), f"{p.forecast_total - p.approved:+,.0f}" if p.forecast_total != p.approved else "–",
                 str(p.in_service), f"{p.npv:,.1f}", pct(p.irr), str(p.payback_years), p.status.split(";")[0]])
    if p.npv < 0:
        red.update({(i, 6), (i, 7)})
    if p.forecast_total > p.approved:
        red.add((i, 4))
table(s, 0.4, 1.5, 12.5, rows, [0.5, 2.9, 1.7, 1.0, 1.0, 0.9, 0.8, 0.8, 1.0, 1.9], 12, red, left_cols=(0, 1, 2, 9))
text(s, 0.4, 5.3, 12.5, 1.5, [
    f"{', '.join(below.index)} return {', '.join(pct(v) for v in below.irr)} against an {pct(A.WACC)} hurdle: recommend re-scoping (phasing, lower cost) before releasing FY2027 funds.",
    "G1 (Southeast quarry) has the highest value but a permitting delay; $9M of FY2026 spend moved into FY2027.",
], 14)

# 9. Scenarios
s = slide("Scenarios: the plan still generates cash in a downturn",
          "Downside: volume -4 / -3 / -1 pts in FY2027-29 and price -1 pt a year. Upside: volume +1.5 to +2 pts, price +0.5 pt")
rows = [["$M", "Downside", "Base", "Upside"]]
for label, key, f in [("FY2027 budget revenue", "budget_2027_revenue", m), ("FY2027 budget EBITDA", "budget_2027_ebitda", m),
                      ("FY2031 revenue", "revenue_2031", m), ("FY2031 EBITDA", "ebitda_2031", m), ("FY2031 EBITDA margin", "margin_2031", pct),
                      ("EBITDA CAGR FY26-31", "ebitda_cagr", pct), ("Five-year capex", "capex_5yr", m),
                      ("Five-year free cash flow", "fcf_5yr", m)]:
    rows.append([label] + [f(scen[k][key]) for k in (2, 1, 3)])
table(s, 0.5, 1.6, 7.6, rows, [3.1, 1.5, 1.5, 1.5], 14)
text(s, 8.5, 1.6, 4.4, 5.0, [
    f"In the downside, FY2031 EBITDA is ${m(base['ebitda_2031'] - scen[2]['ebitda_2031'])}M below base and slightly below FY2026.",
    "Maintenance capex scales with revenue, so capex falls with volume and cash conversion holds.",
    "Growth capex is unchanged across scenarios; G4, G6 and G8 are the first projects to defer if the downside develops.",
], 15)

# 10. Risks and decisions
s = slide("Key risks and decisions requested")
text(s, 0.6, 1.5, 6.0, 5.3, ["Key risks", "Federal highway funding: IIJA surface transportation funding was authorized through FY2026 (September 2026); the size and timing of the next authorization drive aggregates and asphalt volumes.",
                             f"Housing affordability: the 30-year mortgage rate is {latest['MORTGAGE30US'][1]:.2f}% (latest week) and housing starts are flat; ready-mix is most exposed.",
                             "Cost inflation: energy, labor and freight above plan would erode the price-cost spread.",
                             "Permitting: greenfield quarry timing remains uncertain."], 15)
text(s, 7.0, 1.5, 5.9, 5.3, ["Decisions requested", f"Approve the FY2027 budget: revenue ${m(c.loc[2027, 'revenue'])}M, EBITDA ${m(c.loc[2027, 'ebitda'])}M.",
                             f"Approve the FY2027 capex envelope of ${m(c.loc[2027, 'capex'])}M.",
                             "Approve the $6M cost increase on the kiln efficiency project (G3).",
                             "Return G4 and G6 with re-scoped business cases in Q2 FY2027."], 15)
for box in s.shapes:
    if box.has_text_frame and box.text_frame.paragraphs[0].text in ("Key risks", "Decisions requested"):
        p = box.text_frame.paragraphs[0]
        p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, Pt(19)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(OUT)
