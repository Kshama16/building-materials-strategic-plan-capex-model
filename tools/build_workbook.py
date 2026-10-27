"""Builds model/Strategic_Plan_Consolidation_Capex_Model.xlsx with live Excel formulas.

Only inputs (blue) are written as values; every other number is a formula, so changing an
input or the scenario selector recalculates the whole model.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import assumptions as A

OUT = Path(__file__).resolve().parents[1] / "model" / "Strategic_Plan_Consolidation_Capex_Model.xlsx"

BLUE = Font(color="0000CC")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1F3864")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_BORDER = Border(top=Side(style="thin"))
M = '#,##0.0;(#,##0.0);"–"'      # $M
P = '0.0%;(0.0%);"–"'
N2 = '#,##0.00'

YEAR_COL = {y: 3 + i for i, y in enumerate(A.YEARS)}           # 2025 -> C ... 2031 -> I
YEAR_LABEL = {2025: "FY2025A", 2026: "FY2026F (9+3)", 2027: "FY2027 Budget"}
OPCO_NAMES = [o["name"] for o in A.OPCOS]


def yl(y):
    return L(YEAR_COL[y])


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def header(ws, row, labels, start_col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, start_col + i, t)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def year_header(ws, row, label="$ millions"):
    header(ws, row, [label, ""] + [YEAR_LABEL.get(y, f"FY{y} Plan") for y in A.YEARS])
    for y in A.YEARS:                       # numeric year above the label row, used by formulas
        ws.cell(row - 1, YEAR_COL[y], y).font = Font(color="808080", size=8)


def put(ws, row, col, value, fmt=None, font=None):
    c = ws.cell(row, col, value)
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    return c


# ----------------------------------------------------------------------------------------
def cover(wb):
    ws = wb.active
    ws.title = "Cover"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = f"{A.COMPANY}: Five-Year Strategic Plan, Consolidation & CAPEX Model"
    ws["A1"].font = TITLE
    ws["A2"] = "Illustrative company and figures (fictional). Macro data is real, from FRED (Federal Reserve Bank of St. Louis)."
    ws["A2"].font = Font(italic=True, color="C00000")
    rows = [
        ("Purpose", "Consolidates four operating companies into the FY2027 annual budget and FY2027-FY2031 strategic plan, tracks FY2026 CAPEX against budget, evaluates growth projects and builds free cash flow."),
        ("Scenario", "Change Inputs!B3 (1 = Base, 2 = Downside, 3 = Upside). Every sheet recalculates."),
        ("Update the forecast", "Enter new FY2026 year-to-date CAPEX on Inputs and Projects, refresh macro data on Macro (FY2026 driver growth links to the year-to-date change)."),
        ("Conventions", "Blue = input. Black = formula. $ millions unless stated. Capex variance: + = under budget."),
        ("", ""),
        ("Sheets", ""),
        ("Inputs", "Scenario selector, group assumptions, operating company drivers, price growth, scenario adjustments"),
        ("Macro", "FRED housing starts, construction spending, highway spending, mortgage rate, materials PPI; plan driver growth"),
        ("Projects", "Growth CAPEX register: approved vs. forecast, FY2026 budget vs. forecast, status, NPV / IRR / payback"),
        ("OpCo_Plan", "Volume x price build by operating company, EBITDA, maintenance and growth CAPEX"),
        ("Consolidation", "Group P&L with intercompany eliminations, D&A and PP&E roll-forward, taxes, working capital, free cash flow"),
        ("Capex_Report", "FY2026 CAPEX budget vs. YTD vs. forecast by operating company with flags; five-year CAPEX plan"),
        ("Board_Summary", "Headline KPIs for the active scenario and the FY2026-FY2031 EBITDA bridge"),
        ("Checks", "Integrity checks rolled up to one ALL CHECKS PASS cell"),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ws.cell(i, 1, a).font = BOLD
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True)


INP = {}   # cell references on Inputs, filled by inputs()


def inputs(wb):
    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Inputs"
    ws["A1"].font = TITLE
    ws["A3"] = "Active scenario (1 Base, 2 Downside, 3 Upside)"
    put(ws, 3, 2, A.SCENARIO, font=BLUE)
    ws["C3"] = '=CHOOSE(Scenario,"Base","Downside","Upside")'
    ws["C3"].font = BOLD
    name(wb, "Scenario", "Inputs!$B$3")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3")
    ws.add_data_validation(dv)
    dv.add("B3")

    scalars = [("CorpCost2025", "Corporate costs FY2025A ($M)", A.CORP_COST_2025, M),
               ("CorpInfl", "Corporate cost inflation", A.CORP_INFL, P),
               ("TaxRate", "Cash tax rate", A.TAX_RATE, P),
               ("NWCPct", "Net working capital, % of revenue", A.NWC_PCT, P),
               ("PPE2025", "Net PP&E, end FY2025 ($M)", A.PPE_2025, M),
               ("DA2025", "D&A FY2025A ($M)", A.DA_2025, M),
               ("AssetLife", "Average asset life (years); D&A = opening PP&E / life", A.ASSET_LIFE, "0"),
               ("WACC", "WACC / hurdle rate for growth projects", A.WACC, P),
               ("VarPct", "CAPEX variance flag: % threshold", A.VAR_PCT, P),
               ("VarAmt", "CAPEX variance flag: $M threshold", A.VAR_AMT, M)]
    ws["A5"] = "Group assumptions"
    ws["A5"].font = BOLD
    for i, (nm, label, val, fmt) in enumerate(scalars, start=6):
        ws.cell(i, 1, label)
        put(ws, i, 2, val, fmt, BLUE)
        name(wb, nm, f"Inputs!$B${i}")

    r0 = 19
    ws.cell(r0 - 1, 1, "Operating company drivers (FY2025A base)").font = BOLD
    cols = ["Operating company", "Unit", "Volume driver", "Elasticity to driver", "Base volume growth",
            "Volume FY2025A (M)", "Price / unit FY2025A ($)", "Variable cost / unit FY2025A ($)",
            "Fixed costs FY2025A ($M)", "Variable cost inflation", "Fixed cost inflation",
            "Maintenance capex % revenue", "Capex FY2025A ($M)", "Intercompany sales % revenue",
            "FY2026 maint. capex budget ($M)", "FY2026 maint. capex YTD Sep ($M)"]
    header(ws, r0, cols)
    fmts = [None, None, None, "0.00", P, N2, N2, N2, M, P, P, P, M, P, M, M]
    for i, o in enumerate(A.OPCOS):
        r = r0 + 1 + i
        m26 = A.MAINT_2026[o["name"]]
        vals = [o["name"], o["unit"], o["driver"], o["elasticity"], o["base_growth"], o["volume"], o["price"],
                o["var_cost"], o["fixed"], o["var_infl"], o["fixed_infl"], o["maint_pct"], o["capex_2025"],
                o["ic_pct"], m26["budget"], m26["ytd"]]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            put(ws, r, j, v, f, BLUE if j > 1 else BOLD)
        INP[o["name"]] = r

    pr0 = 27
    ws.cell(pr0 - 1, 1, "Price growth by year").font = BOLD
    header(ws, pr0, ["Operating company", ""] + [""] + [f"FY{y}" for y in A.PLAN_YEARS])
    for i, o in enumerate(A.OPCOS):
        r = pr0 + 1 + i
        ws.cell(r, 1, o["name"])
        for y in A.PLAN_YEARS:
            put(ws, r, YEAR_COL[y], A.PRICE_GROWTH[o["name"]][y], P, BLUE)
        INP["price_" + o["name"]] = r

    s0 = 35
    ws.cell(s0 - 1, 1, "Scenario adjustments (plan years FY2027-FY2031)").font = BOLD
    header(ws, s0, ["Volume growth adjustment", ""] + [""] + [f"FY{y}" for y in A.PLAN_YEARS])
    for k, sname in A.SCENARIOS.items():
        r = s0 + k
        ws.cell(r, 1, sname)
        for y in A.PLAN_YEARS:
            put(ws, r, YEAR_COL[y], A.SCEN_VOL_ADJ[k].get(y, 0.0), P, BLUE)
    r = s0 + 4
    ws.cell(r, 1, "Active").font = BOLD
    for y in A.PLAN_YEARS:
        c = yl(y)
        put(ws, r, YEAR_COL[y], f"=CHOOSE(Scenario,{c}{s0 + 1},{c}{s0 + 2},{c}{s0 + 3})", P, BOLD)
    INP["vol_adj"] = r

    p0 = 42
    ws.cell(p0 - 1, 1, "Price growth adjustment (added each plan year)").font = BOLD
    for k, sname in A.SCENARIOS.items():
        ws.cell(p0 + k - 1, 1, sname)
        put(ws, p0 + k - 1, 2, A.SCEN_PRICE_ADJ[k], P, BLUE)
    r = p0 + 3
    ws.cell(r, 1, "Active, by year").font = BOLD
    put(ws, r, YEAR_COL[2026], 0, P)
    for y in range(2027, 2032):
        put(ws, r, YEAR_COL[y], f"=CHOOSE(Scenario,$B${p0},$B${p0 + 1},$B${p0 + 2})", P, BOLD)
    INP["price_adj"] = r

    ws.column_dimensions["A"].width = 44
    for c in range(2, 17):
        ws.column_dimensions[L(c)].width = 14
    ws.row_dimensions[r0].height = 45
    ws.freeze_panes = "B4"


MAC = {}


def macro(wb):
    ws = wb.create_sheet("Macro")
    hist, ytd, latest = A.macro_annual()
    ws["A1"] = "Macro backdrop: US construction demand (source: FRED, Federal Reserve Bank of St. Louis)"
    ws["A1"].font = TITLE
    ws["A2"] = "Annual averages of monthly (weekly for mortgage rate) data; FY2026 = year to date."
    hist_years = list(hist.index)
    header(ws, 4, ["Series", "FRED ID"] + [str(y) for y in hist_years] + ["2026 YTD vs. 2025 YTD", "Latest value", "Latest date"])
    for i, sid in enumerate(A.MACRO_SERIES):
        r = 5 + i
        ws.cell(r, 1, A.MACRO_SERIES[sid])
        ws.cell(r, 2, sid)
        for j, y in enumerate(hist_years):
            put(ws, r, 3 + j, round(float(hist.loc[y, sid]), 2), '#,##0.00' if hist.loc[y, sid] < 100 else '#,##0', BLUE)
        k = 3 + len(hist_years)
        put(ws, r, k, round(ytd[sid], 4), P, BLUE)
        put(ws, r, k + 1, latest[sid][1], '#,##0.00', BLUE)
        put(ws, r, k + 2, latest[sid][0], None, BLUE)
        MAC["ytd_" + sid] = f"${L(k)}${r}"

    ws.cell(12, 1, "Volume driver growth used in the plan").font = BOLD
    ws.cell(13, 1, "FY2026 links to the actual year-to-date change above; FY2027-FY2031 are planning assumptions.")
    year_header(ws, 15, "Driver")
    for i, sid in enumerate(A.DRIVERS):
        r = 16 + i
        ws.cell(r, 1, A.MACRO_SERIES[sid])
        ws.cell(r, 2, sid)
        put(ws, r, YEAR_COL[2026], f"={MAC['ytd_' + sid]}", P)
        for y in range(2027, 2032):
            put(ws, r, YEAR_COL[y], A.MACRO_PLAN[sid][y], P, BLUE)
        MAC[sid] = r
    ws.column_dimensions["A"].width = 50
    for c in range(2, 15):
        ws.column_dimensions[L(c)].width = 13


PRJ = {"first": 5}


def projects(wb):
    ws = wb.create_sheet("Projects")
    ws["A1"] = "Growth CAPEX project register"
    ws["A1"].font = TITLE
    ws["A2"] = "Spend schedule = latest forecast by year. Variance: + = under budget. Returns use after-tax EBITDA uplift from the in-service year, 20-year horizon."
    spend_years = list(range(2026, 2032))
    cols = (["ID", "Operating company", "Project", "Type", "Approved budget"] + [f"Spend FY{y}" for y in spend_years]
            + ["Forecast total", "Over/(under) approved", "FY2026 budget", "FY2026 YTD Sep", "FY2026 forecast",
               "FY2026 variance", "In-service year", "Annual EBITDA uplift", "NPV @ WACC", "IRR", "Payback (years)", "Status"])
    header(ws, 4, cols)
    first = PRJ["first"]
    last = first + len(A.PROJECTS) - 1
    PRJ["last"] = last
    spend_col = {y: 6 + i for i, y in enumerate(spend_years)}           # F..K
    PRJ["spend_col"] = spend_col
    cf0, cum0 = last + 5, last + 5 + len(A.PROJECTS) + 3
    for i, p in enumerate(A.PROJECTS):
        r = first + i
        for j, v in enumerate([p["id"], p["opco"], p["name"], "Growth"], start=1):
            ws.cell(r, j, v)
        put(ws, r, 5, p["approved"], M, BLUE)
        for y, c in spend_col.items():
            put(ws, r, c, p["spend"].get(y, 0.0), M, BLUE)
        put(ws, r, 12, f"=SUM(F{r}:K{r})", M)
        put(ws, r, 13, f"=L{r}-E{r}", M)
        put(ws, r, 14, p["fy26_budget"], M, BLUE)
        put(ws, r, 15, p["fy26_ytd"], M, BLUE)
        put(ws, r, 16, f"=F{r}", M)
        put(ws, r, 17, f"=N{r}-P{r}", M)
        put(ws, r, 18, p["in_service"], "0", BLUE)
        put(ws, r, 19, p["uplift"], M, BLUE)
        cfr, cumr = cf0 + 1 + i, cum0 + 1 + i
        put(ws, r, 20, f"=C{cfr}+NPV(WACC,D{cfr}:V{cfr})", M)
        put(ws, r, 21, f"=IRR(C{cfr}:V{cfr})", P)
        put(ws, r, 22, f'=COUNTIF(C{cumr}:V{cumr},"<0")', "0")
        put(ws, r, 23, p["status"], None, BLUE)
    tr = last + 1
    ws.cell(tr, 3, "Total").font = BOLD
    for c in [5] + list(spend_col.values()) + [12, 13, 14, 15, 16, 17, 19, 20]:
        put(ws, tr, c, f"=SUM({L(c)}{first}:{L(c)}{last})", M, BOLD).border = TOTAL_BORDER
    ws.conditional_formatting.add(f"T{first}:T{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(f"M{first}:M{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="C00000", bold=True)))

    # After-tax project cash flows
    ws.cell(cf0 - 1, 1, "Project cash flows: -spend + EBITDA uplift x (1 - tax) once in service").font = BOLD
    header(ws, cf0, ["ID", "Project"] + [str(y) for y in A.RETURN_YEARS])
    header(ws, cum0, ["ID", "Cumulative"] + [str(y) for y in A.RETURN_YEARS])
    for i in range(len(A.PROJECTS)):
        r, cfr, cumr = first + i, cf0 + 1 + i, cum0 + 1 + i
        ws.cell(cfr, 1, f"=A{r}")
        ws.cell(cfr, 2, f"=C{r}")
        ws.cell(cumr, 1, f"=A{r}")
        for j, y in enumerate(A.RETURN_YEARS):
            c = 3 + j
            spend = f"-{L(spend_col[y])}{r}" if y in spend_col else "0"
            put(ws, cfr, c, f"={spend}+IF({y}>=$R${r},$S${r}*(1-TaxRate),0)", M)
            prev = f"{L(c - 1)}{cumr}+" if j else ""
            put(ws, cumr, c, f"={prev}{L(c)}{cfr}", M)
    widths = [5, 20, 34, 9, 11] + [10] * 6 + [11, 11, 11, 11, 11, 11, 10, 11, 11, 8, 9, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "D5"


OP = {}
ROWS = ["Macro driver growth", "Volume growth", "Volume (M units)", "Price growth", "Price per unit ($)",
        "Revenue", "Variable cost per unit ($)", "Variable costs", "Fixed costs", "Growth project EBITDA",
        "EBITDA", "EBITDA margin", "Maintenance capex", "Growth capex", "Total capex"]


def opco_plan(wb):
    ws = wb.create_sheet("OpCo_Plan")
    ws["A1"] = "Operating company plans: volume x price build"
    ws["A1"].font = TITLE
    year_header(ws, 3)
    first, last = PRJ["first"], PRJ["last"]
    for b, o in enumerate(A.OPCOS):
        n = o["name"]
        top = 4 + b * 17
        ir, pr = INP[n], INP["price_" + n]
        rr = {k: top + 1 + i for i, k in enumerate(["drv", "vg", "vol", "pg", "price", "rev", "vc", "var", "fixed",
                                                      "grow", "ebitda", "margin", "maint", "gcap", "capex"])}
        OP[n] = rr
        c = ws.cell(top, 1, f"{n} (driver: {o['driver']}, volume in million {o['unit']})")
        c.font, c.fill = BOLD, SUB_FILL
        for col in range(2, 10):
            ws.cell(top, col).fill = SUB_FILL
        for k, label in zip(rr, ROWS):
            ws.cell(rr[k], 1, label)
        C = yl(2025)
        put(ws, rr["vol"], 3, f"=Inputs!$F${ir}", N2)
        put(ws, rr["price"], 3, f"=Inputs!$G${ir}", N2)
        put(ws, rr["vc"], 3, f"=Inputs!$H${ir}", N2)
        put(ws, rr["fixed"], 3, f"=Inputs!$I${ir}", M)
        put(ws, rr["capex"], 3, f"=Inputs!$M${ir}", M)
        for y in A.YEARS:
            c, prev = yl(y), (yl(y - 1) if y > 2025 else None)
            if y > 2025:
                put(ws, rr["drv"], YEAR_COL[y], f"=Macro!{c}${MAC[o['driver']]}", P)
                put(ws, rr["vg"], YEAR_COL[y], f"=Inputs!$E${ir}+Inputs!$D${ir}*{c}{rr['drv']}+Inputs!{c}${INP['vol_adj']}", P)
                put(ws, rr["vol"], YEAR_COL[y], f"={prev}{rr['vol']}*(1+{c}{rr['vg']})", N2)
                put(ws, rr["pg"], YEAR_COL[y], f"=Inputs!{c}${pr}+Inputs!{c}${INP['price_adj']}", P)
                put(ws, rr["price"], YEAR_COL[y], f"={prev}{rr['price']}*(1+{c}{rr['pg']})", N2)
                put(ws, rr["vc"], YEAR_COL[y], f"={prev}{rr['vc']}*(1+Inputs!$J${ir})", N2)
                put(ws, rr["fixed"], YEAR_COL[y], f"={prev}{rr['fixed']}*(1+Inputs!$K${ir})", M)
                put(ws, rr["maint"], YEAR_COL[y], f"=Inputs!$L${ir}*{c}{rr['rev']}", M)
                sc = L(PRJ["spend_col"][y])
                put(ws, rr["gcap"], YEAR_COL[y], f'=SUMIF(Projects!$B${first}:$B${last},"{n}",Projects!${sc}${first}:${sc}${last})', M)
                put(ws, rr["capex"], YEAR_COL[y], f"={c}{rr['maint']}+{c}{rr['gcap']}", M)
            put(ws, rr["rev"], YEAR_COL[y], f"={c}{rr['vol']}*{c}{rr['price']}", M)
            put(ws, rr["var"], YEAR_COL[y], f"={c}{rr['vol']}*{c}{rr['vc']}", M)
            put(ws, rr["grow"], YEAR_COL[y],
                f'=SUMIFS(Projects!$S${first}:$S${last},Projects!$B${first}:$B${last},"{n}",Projects!$R${first}:$R${last},"<="&{c}$2)', M)
            put(ws, rr["ebitda"], YEAR_COL[y], f"={c}{rr['rev']}-{c}{rr['var']}-{c}{rr['fixed']}+{c}{rr['grow']}", M, BOLD)
            put(ws, rr["margin"], YEAR_COL[y], f"={c}{rr['ebitda']}/{c}{rr['rev']}", P)
        for k in ("rev", "ebitda", "capex"):
            ws.cell(rr[k], 1).font = BOLD
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 2
    for y in A.YEARS:
        ws.column_dimensions[yl(y)].width = 14
    ws.freeze_panes = "C4"


CON = {}


def consolidation(wb):
    ws = wb.create_sheet("Consolidation")
    ws["A1"] = f"{A.COMPANY}: consolidated plan"
    ws["A1"].font = TITLE
    ws["C1"] = '="Scenario: "&Inputs!$C$3'
    ws["C1"].font = BOLD
    year_header(ws, 4)
    r = 5
    lines = []

    def line(key, label, fn, fmt=M, bold=False, total=False, years=A.YEARS):
        nonlocal r
        CON[key] = r
        ws.cell(r, 1, label).font = BOLD if bold else Font()
        for y in years:
            v = fn(y, yl(y))
            if v is not None:
                cell = put(ws, r, YEAR_COL[y], v, fmt, BOLD if bold else None)
                if total:
                    cell.border = TOTAL_BORDER
        r += 1

    def blank():
        nonlocal r
        r += 1

    for n in OPCO_NAMES:
        line("rev_" + n, f"  Revenue: {n}", lambda y, c, n=n: f"=OpCo_Plan!{c}{OP[n]['rev']}")
    line("opco_rev", "Total operating company revenue", lambda y, c: f"=SUM({c}{CON['rev_' + OPCO_NAMES[0]]}:{c}{CON['rev_' + OPCO_NAMES[-1]]})", total=True)
    line("ic", "Intercompany eliminations",
         lambda y, c: "=-(" + "+".join(f"{c}{CON['rev_' + o['name']]}*Inputs!$N${INP[o['name']]}" for o in A.OPCOS) + ")")
    line("rev", "Consolidated revenue", lambda y, c: f"={c}{CON['opco_rev']}+{c}{CON['ic']}", bold=True, total=True)
    blank()
    for n in OPCO_NAMES:
        line("ebitda_" + n, f"  EBITDA: {n}", lambda y, c, n=n: f"=OpCo_Plan!{c}{OP[n]['ebitda']}")
    line("corp", "Corporate costs", lambda y, c: "=-CorpCost2025" if y == 2025 else f"={yl(y - 1)}{CON['corp']}*(1+CorpInfl)")
    line("ebitda", "Consolidated EBITDA",
         lambda y, c: f"=SUM({c}{CON['ebitda_' + OPCO_NAMES[0]]}:{c}{CON['corp']})", bold=True, total=True)
    line("margin", "EBITDA margin", lambda y, c: f"={c}{CON['ebitda']}/{c}{CON['rev']}", fmt=P)
    line("ic_note", "Intercompany sales are eliminated at cost, so EBITDA is unaffected", lambda y, c: None)
    blank()
    ppe_row = r + 7  # computed below; D&A references prior-year PP&E
    line("da", "Depreciation & amortization",
         lambda y, c: "=DA2025" if y == 2025 else f"={yl(y - 1)}{ppe_row}/AssetLife")
    line("ebit", "EBIT", lambda y, c: f"={c}{CON['ebitda']}-{c}{CON['da']}", total=True)
    line("tax", "Cash taxes", lambda y, c: None if y == 2025 else f"={c}{CON['ebit']}*TaxRate")
    blank()
    line("maint", "Maintenance capex", lambda y, c: None if y == 2025 else
         "=" + "+".join(f"OpCo_Plan!{c}{OP[n]['maint']}" for n in OPCO_NAMES))
    line("gcap", "Growth capex", lambda y, c: None if y == 2025 else
         "=" + "+".join(f"OpCo_Plan!{c}{OP[n]['gcap']}" for n in OPCO_NAMES))
    line("capex", "Total capex", lambda y, c: "=" + "+".join(f"OpCo_Plan!{c}{OP[n]['capex']}" for n in OPCO_NAMES)
         if y == 2025 else f"={c}{CON['maint']}+{c}{CON['gcap']}", bold=True, total=True)
    assert r == ppe_row, (r, ppe_row)
    line("ppe", "Net PP&E (closing)", lambda y, c: "=PPE2025" if y == 2025 else
         f"={yl(y - 1)}{CON['ppe']}+{c}{CON['capex']}-{c}{CON['da']}")
    line("capex_da", "Capex / D&A", lambda y, c: f"={c}{CON['capex']}/{c}{CON['da']}", fmt='0.00"x"')
    blank()
    line("nwc", "Net working capital", lambda y, c: f"={c}{CON['rev']}*NWCPct")
    line("dnwc", "Increase in net working capital", lambda y, c: None if y == 2025 else f"={c}{CON['nwc']}-{yl(y - 1)}{CON['nwc']}")
    blank()
    line("fcf", "Free cash flow (EBITDA - taxes - capex - increase in NWC)", lambda y, c: None if y == 2025 else
         f"={c}{CON['ebitda']}-{c}{CON['tax']}-{c}{CON['capex']}-{c}{CON['dnwc']}", bold=True, total=True)
    line("fcf_conv", "FCF conversion (FCF / EBITDA)", lambda y, c: None if y == 2025 else f"={c}{CON['fcf']}/{c}{CON['ebitda']}", fmt=P)
    ws.cell(CON["ic_note"], 1).font = Font(italic=True, color="808080")
    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 2
    for y in A.YEARS:
        ws.column_dimensions[yl(y)].width = 14
    ws.freeze_panes = "C5"


CAP = {}


def capex_report(wb):
    ws = wb.create_sheet("Capex_Report")
    ws["A1"] = "FY2026 CAPEX: budget vs. year to date vs. forecast"
    ws["A1"].font = TITLE
    ws["A2"] = '="Flag when variance is at least "&TEXT(VarPct,"0%")&" and $"&TEXT(VarAmt,"0")&"M. Variance: + = under budget."'
    cols = ["Operating company", "Maint. budget", "Maint. YTD", "Maint. forecast", "Growth budget", "Growth YTD",
            "Growth forecast", "Total budget", "Total YTD", "Total forecast", "% of forecast spent", "Variance",
            "Variance %", "Flag"]
    header(ws, 4, cols)
    first, last = PRJ["first"], PRJ["last"]
    D = yl(2026)
    for i, n in enumerate(OPCO_NAMES):
        r = 5 + i
        ws.cell(r, 1, n)
        rng = lambda col: f'SUMIF(Projects!$B${first}:$B${last},"{n}",Projects!${col}${first}:${col}${last})'
        vals = [f"=Inputs!$O${INP[n]}", f"=Inputs!$P${INP[n]}", f"=OpCo_Plan!{D}{OP[n]['maint']}",
                f"={rng('N')}", f"={rng('O')}", f"={rng('P')}",
                f"=B{r}+E{r}", f"=C{r}+F{r}", f"=D{r}+G{r}", f"=I{r}/J{r}", f"=H{r}-J{r}", f"=L{r}/H{r}",
                f'=IF(AND(ABS(M{r})>=VarPct,ABS(L{r})>=VarAmt),IF(L{r}>0,"UNDER","OVER"),"")']
        for j, v in enumerate(vals, start=2):
            put(ws, r, j, v, P if j in (11, 13) else (None if j == 14 else M))
    tr = 5 + len(OPCO_NAMES)
    CAP["first"], CAP["total"] = 5, tr
    ws.cell(tr, 1, "Total").font = BOLD
    for j in range(2, 11):
        put(ws, tr, j, f"=SUM({L(j)}5:{L(j)}{tr - 1})", M, BOLD).border = TOTAL_BORDER
    put(ws, tr, 11, f"=I{tr}/J{tr}", P, BOLD)
    put(ws, tr, 12, f"=H{tr}-J{tr}", M, BOLD)
    put(ws, tr, 13, f"=L{tr}/H{tr}", P, BOLD)
    put(ws, tr, 14, f'=IF(AND(ABS(M{tr})>=VarPct,ABS(L{tr})>=VarAmt),IF(L{tr}>0,"UNDER","OVER"),"")', None, BOLD)
    ws.conditional_formatting.add(f"N5:N{tr}", FormulaRule(formula=['N5<>""'], font=Font(color="C00000", bold=True)))

    r0 = tr + 3
    ws.cell(r0 - 1, 1, "Five-year CAPEX plan").font = BOLD
    plan_years = list(range(2027, 2032))
    header(ws, r0, ["$ millions"] + [f"FY{y}" for y in plan_years] + ["5-year total", "% of total"])
    r = r0 + 1
    rows = [(f"{n}", lambda c, n=n: f"=OpCo_Plan!{c}{OP[n]['capex']}") for n in OPCO_NAMES]
    rows += [("Total capex", lambda c: f"=Consolidation!{c}{CON['capex']}"),
             ("  of which maintenance", lambda c: f"=Consolidation!{c}{CON['maint']}"),
             ("  of which growth", lambda c: f"=Consolidation!{c}{CON['gcap']}"),
             ("Capex / D&A", lambda c: f"=Consolidation!{c}{CON['capex_da']}"),
             ("Capex % of revenue", lambda c: f"=Consolidation!{c}{CON['capex']}/Consolidation!{c}{CON['rev']}")]
    total_row = r0 + 1 + len(OPCO_NAMES)
    for label, fn in rows:
        ws.cell(r, 1, label).font = BOLD if label == "Total capex" else Font()
        ratio = label.startswith("Capex")
        for j, y in enumerate(plan_years):
            put(ws, r, 2 + j, fn(yl(y)), '0.00"x"' if label == "Capex / D&A" else (P if ratio else M))
        if not ratio:
            put(ws, r, 7, f"=SUM(B{r}:F{r})", M, BOLD)
            put(ws, r, 8, f"=G{r}/$G${total_row}", P)
        CAP[label.strip()] = r
        r += 1
    ws.column_dimensions["A"].width = 26
    for j in range(2, 15):
        ws.column_dimensions[L(j)].width = 12
    ws.row_dimensions[4].height = 32


BRD = {}


def board_summary(wb):
    ws = wb.create_sheet("Board_Summary")
    ws["A1"] = f"{A.COMPANY}: plan summary for the board"
    ws["A1"].font = TITLE
    ws["A2"] = '="Scenario: "&Inputs!$C$3'
    ws["A2"].font = BOLD
    header(ws, 4, ["$ millions", "FY2025A", "FY2026F", "FY2027 Budget", "FY2031 Plan", "CAGR FY26-31"])
    kpis = [("Revenue", "rev", M), ("EBITDA", "ebitda", M), ("EBITDA margin", "margin", P),
            ("Total capex", "capex", M), ("Free cash flow", "fcf", M)]
    for i, (label, key, fmt) in enumerate(kpis):
        r = 5 + i
        ws.cell(r, 1, label).font = BOLD
        for j, y in enumerate([2025, 2026, 2027, 2031]):
            if key == "fcf" and y == 2025:
                continue
            put(ws, r, 2 + j, f"=Consolidation!{yl(y)}{CON[key]}", fmt)
        if key in ("rev", "ebitda"):
            put(ws, r, 6, f"=(E{r}/C{r})^(1/5)-1", P)
        BRD[key] = r
    r = 11
    agg = [("5-year capex FY27-31", f"=SUM(Consolidation!{yl(2027)}{CON['capex']}:Consolidation!{yl(2031)}{CON['capex']})", M),
           ("  of which growth", f"=SUM(Consolidation!{yl(2027)}{CON['gcap']}:Consolidation!{yl(2031)}{CON['gcap']})", M),
           ("5-year free cash flow FY27-31", f"=SUM(Consolidation!{yl(2027)}{CON['fcf']}:Consolidation!{yl(2031)}{CON['fcf']})", M),
           ("5-year FCF conversion", None, P),
           ("Growth projects below hurdle (NPV < 0)", f'=COUNTIF(Projects!$T${PRJ["first"]}:$T${PRJ["last"]},"<0")', "0"),
           ("FY2026 capex forecast vs. budget (+ = under)", f"=Capex_Report!L{CAP['total']}", M)]
    for label, f, fmt in agg:
        ws.cell(r, 1, label).font = BOLD
        if label == "5-year FCF conversion":
            f = f"=B{r - 1}/SUM(Consolidation!{yl(2027)}{CON['ebitda']}:Consolidation!{yl(2031)}{CON['ebitda']})"
        put(ws, r, 2, f, fmt)
        BRD[label] = r
        r += 1

    b0 = r + 2
    ws.cell(b0 - 1, 1, "EBITDA bridge FY2026F to FY2031 Plan").font = BOLD
    header(ws, b0, ["Driver", "$ millions"])
    s, e = yl(2026), yl(2031)
    terms = {
        "volume": lambda rr: f"({e}{rr['vol']}-{s}{rr['vol']})*({s}{rr['price']}-{s}{rr['vc']})",
        "price": lambda rr: f"{e}{rr['vol']}*({e}{rr['price']}-{s}{rr['price']})",
        "var_cost": lambda rr: f"-{e}{rr['vol']}*({e}{rr['vc']}-{s}{rr['vc']})",
        "fixed": lambda rr: f"-({e}{rr['fixed']}-{s}{rr['fixed']})",
        "growth": lambda rr: f"({e}{rr['grow']}-{s}{rr['grow']})",
    }
    bridge = [("FY2026F EBITDA", "start", f"=Consolidation!{s}{CON['ebitda']}"),
              ("Volume", "volume", None), ("Price", "price", None),
              ("Variable cost inflation", "var_cost", None), ("Fixed cost inflation", "fixed", None),
              ("Growth projects in service", "growth", None),
              ("Corporate costs", "corporate", f"=Consolidation!{e}{CON['corp']}-Consolidation!{s}{CON['corp']}"),
              ("FY2031 Plan EBITDA", "end", None)]
    for i, (label, key, f) in enumerate(bridge):
        r = b0 + 1 + i
        ws.cell(r, 1, label).font = BOLD if key in ("start", "end") else Font()
        if f is None and key != "end":
            f = "=" + "+".join(_qualify(terms[key](OP[n]), s, e) for n in OPCO_NAMES)
        if key == "end":
            f = f"=SUM(B{b0 + 1}:B{r - 1})"
        put(ws, r, 2, f, M, BOLD if key in ("start", "end") else None)
        BRD["bridge_" + key] = r
    ws.column_dimensions["A"].width = 46
    for j in range(2, 7):
        ws.column_dimensions[L(j)].width = 15


def _qualify(expr, s, e):
    """Prefixes OpCo_Plan! to every FY2026 / FY2031 cell reference in a bridge term."""
    import re
    return re.sub(rf"\b([{s}{e}])(\d+)\b", r"OpCo_Plan!\1\2", expr)


def checks(wb):
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Integrity checks"
    ws["A1"].font = TITLE
    header(ws, 3, ["Check", "Pass?", "Difference"])
    c = lambda key, y: f"Consolidation!{yl(y)}{CON[key]}"
    rng = lambda key, y0, y1: f"Consolidation!{yl(y0)}{CON[key]}:{yl(y1)}{CON[key]}"
    opco_rev = "+".join(f"SUM(OpCo_Plan!{yl(2025)}{OP[n]['rev']}:{yl(2031)}{OP[n]['rev']})" for n in OPCO_NAMES)
    ic = "+".join(f"SUM(OpCo_Plan!{yl(2025)}{OP[o['name']]['rev']}:{yl(2031)}{OP[o['name']]['rev']})*Inputs!$N${INP[o['name']]}" for o in A.OPCOS)
    first, last = PRJ["first"], PRJ["last"]
    items = [
        ("Consolidated revenue = operating company revenue - intercompany (all years)",
         f"SUM({rng('rev', 2025, 2031)})-(({opco_rev})-({ic}))"),
        ("Consolidated EBITDA = sum of operating company EBITDA + corporate (all years)",
         f"SUM({rng('ebitda', 2025, 2031)})-(" + "+".join(f"SUM(OpCo_Plan!{yl(2025)}{OP[n]['ebitda']}:{yl(2031)}{OP[n]['ebitda']})" for n in OPCO_NAMES) + f"+SUM({rng('corp', 2025, 2031)}))"),
        ("Total capex = maintenance + growth (FY2026-FY2031)",
         f"SUM({rng('capex', 2026, 2031)})-SUM({rng('maint', 2026, 2031)})-SUM({rng('gcap', 2026, 2031)})"),
        ("Growth capex in plan = project register spend schedule",
         f"SUM({rng('gcap', 2026, 2031)})-SUM(Projects!$F${first}:$K${last})"),
        ("PP&E roll-forward: FY2031 = FY2025 + capex - D&A",
         f"{c('ppe', 2031)}-(PPE2025+SUM({rng('capex', 2026, 2031)})-SUM({rng('da', 2026, 2031)}))"),
        ("Free cash flow = EBITDA - taxes - capex - increase in NWC (FY2026-FY2031)",
         f"SUM({rng('fcf', 2026, 2031)})-(SUM({rng('ebitda', 2026, 2031)})-SUM({rng('tax', 2026, 2031)})-SUM({rng('capex', 2026, 2031)})-SUM({rng('dnwc', 2026, 2031)}))"),
        ("EBITDA bridge ends at FY2031 Plan EBITDA",
         f"Board_Summary!B{BRD['bridge_end']}-{c('ebitda', 2031)}"),
        ("FY2026 capex report forecast = consolidated FY2026 capex",
         f"Capex_Report!J{CAP['total']}-{c('capex', 2026)}"),
        ("Project forecast totals = sum of yearly schedule",
         f"SUM(Projects!$L${first}:$L${last})-SUM(Projects!$F${first}:$K${last})"),
    ]
    for i, (label, diff) in enumerate(items):
        r = 4 + i
        ws.cell(r, 1, label)
        put(ws, r, 3, f"={diff}", '0.000000')
        ws.cell(r, 2, f"=ABS(C{r})<0.001")
    r = 4 + len(items)
    ws.cell(r, 1, "Scenario selector is 1, 2 or 3")
    ws.cell(r, 2, "=OR(Scenario=1,Scenario=2,Scenario=3)")
    r += 1
    ws.cell(r, 1, "ALL CHECKS PASS").font = BOLD
    ws.cell(r, 2, f"=AND(B4:B{r - 1})").font = BOLD
    ws.conditional_formatting.add(f"B4:B{r}", CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"B4:B{r}", CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.column_dimensions["A"].width = 78
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14


def build(path=OUT, scenario=None):
    if scenario is not None:
        A.SCENARIO = scenario
    wb = Workbook()
    cover(wb)
    inputs(wb)
    macro(wb)
    projects(wb)
    opco_plan(wb)
    consolidation(wb)
    capex_report(wb)
    board_summary(wb)
    checks(wb)
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
